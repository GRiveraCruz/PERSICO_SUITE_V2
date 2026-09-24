"""
consignacion.py — Almacén de CONSIGNACIÓN (Almacenes ▸ Consignación).

Mismas funciones que Stock (inventario, ingreso, edición, importación desde
Excel, borrado, exportación) y que Reasignaciones (órdenes que cargan material
a un Job), pero con sus datos en una base separada por cuestiones fiscales:

  • PostgreSQL: tablas `consignacion`, `consignacion_reasignaciones`,
    `consignacion_recuperaciones` y `consignacion_folios`, definidas sobre su
    propio ConsigBase en db.py. Con CONSIG_DATABASE_URL viven en otra base de
    datos; sin ella, en la misma instancia pero en tablas propias.
  • Sin base de datos: archivos JSON en DATA_DIR/CONSIGNACION/.

Diferencias intencionales respecto a Stock (más estrictas, por ser datos fiscales):
  • Cada operación que toca más de una colección (ingreso + recuperación,
    reasignación + inventario) corre en UNA transacción con advisory lock de
    PostgreSQL: o se guarda todo o nada, y dos workers de gunicorn no pueden
    pisarse entre sí.
  • El folio de reasignación (CRA-0000000001) solo se asigna al guardar la
    orden; consultar la lista muestra el siguiente folio sin consumirlo, así que
    la numeración no tiene huecos.
  • Reasignar valida en el servidor que el material exista y que haya
    existencia suficiente (Stock descuenta con max(0, …) y acepta materiales
    inexistentes sin avisar).
"""
import datetime
import io
import json
import unicodedata
from contextlib import contextmanager
from pathlib import Path
from threading import Lock

import openpyxl
from flask import Blueprint, jsonify, request, Response

try:
    import db as _orm
except Exception:          # sin SQLAlchemy instalado → modo JSON
    _orm = None

bp = Blueprint("consignacion", __name__)

FOLIO_PREFIX = "CRA"
FOLIO_WIDTH = 10
ORIGEN = "Consignación"

_lock = Lock()
_is_admin = lambda: False
_data_dir = Path("data") / "CONSIGNACION"

# kind → (archivo JSON, modelo ORM)
_KINDS = {
    "items":    ("consignacion.json",   "Consignacion"),
    "orders":   ("reasignaciones.json", "ConsigReasignacion"),
    "recovery": ("recuperaciones.json", "ConsigRecuperacion"),
}


def setup(app, is_admin_fn, data_dir):
    """Registra el blueprint. Se llama una sola vez desde app.py."""
    global _is_admin, _data_dir
    _is_admin = is_admin_fn
    _data_dir = Path(data_dir) / "CONSIGNACION"
    app.register_blueprint(bp)


def db_enabled():
    return bool(_orm and getattr(_orm, "CONSIG_DB_ENABLED", False))


def _current_user():
    from flask import session
    return session.get("user", "")


def _now():
    return datetime.datetime.now()


# ══════════════════════════════════════════════════════════════════
#  Capa de datos
# ══════════════════════════════════════════════════════════════════
def _json_path(kind):
    return _data_dir / _KINDS[kind][0]


def _json_read(path, default):
    if path.exists():
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return default
    return default


def _json_write(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    tmp.replace(path)          # reemplazo atómico: nunca queda un archivo a medias


def load(kind):
    """Lectura sin candado (para GET y para reportes)."""
    if db_enabled():
        model = getattr(_orm, _KINDS[kind][1])
        s = _orm.get_consig_session()
        try:
            return [r.data for r in s.query(model).order_by(model.id.asc()).all()]
        finally:
            s.close()
    return _json_read(_json_path(kind), [])


def _row_for(kind, rec):
    if kind == "items":
        return _orm.Consignacion(data=rec, content_hash=_orm.canonical_hash(rec), item_id=rec.get("id"))
    if kind == "orders":
        return _orm.ConsigReasignacion(data=rec, order_number=str(rec.get("order_number") or ""))
    return _orm.ConsigRecuperacion(data=rec, content_hash=_orm.canonical_hash(rec), job=rec.get("job"))


class _Tx:
    """Carga/guarda colecciones dentro de una misma transacción."""

    def __init__(self, session=None):
        self.s = session
        self._loaded = {}
        self._dirty = {}
        self._folios = None

    def load(self, kind):
        if kind not in self._loaded:
            if self.s is not None:
                model = getattr(_orm, _KINDS[kind][1])
                self._loaded[kind] = [r.data for r in self.s.query(model).order_by(model.id.asc()).all()]
            else:
                self._loaded[kind] = _json_read(_json_path(kind), [])
        return self._loaded[kind]

    def save(self, kind, records):
        self._loaded[kind] = records
        self._dirty[kind] = records

    def next_folio(self, prefix=FOLIO_PREFIX):
        if self.s is not None:
            from sqlalchemy import text
            n = self.s.execute(text(
                "INSERT INTO consignacion_folios (prefix, n) VALUES (:p, 1) "
                "ON CONFLICT (prefix) DO UPDATE SET n = consignacion_folios.n + 1 RETURNING n"
            ), {"p": prefix}).scalar()
        else:
            if self._folios is None:
                self._folios = _json_read(_data_dir / "folios.json", {})
            n = int(self._folios.get(prefix, 0)) + 1
            self._folios[prefix] = n
        return f"{prefix}-{str(n).zfill(FOLIO_WIDTH)}"

    def flush(self):
        if self.s is not None:
            for kind, records in self._dirty.items():
                model = getattr(_orm, _KINDS[kind][1])
                self.s.query(model).delete()
                for r in records:
                    if kind == "orders" and not r.get("order_number"):
                        continue
                    self.s.add(_row_for(kind, r))
        else:
            for kind, records in self._dirty.items():
                _json_write(_json_path(kind), records)
            if self._folios is not None:
                _json_write(_data_dir / "folios.json", self._folios)


@contextmanager
def tx():
    """Transacción de escritura: todo lo que se guarde dentro se confirma junto,
    o nada si ocurre un error. En PostgreSQL toma un advisory lock propio de
    consignación para serializar escrituras entre workers."""
    if db_enabled():
        from sqlalchemy import text
        s = _orm.get_consig_session()
        try:
            s.execute(text("SELECT pg_advisory_xact_lock(hashtext(:k))"), {"k": "consignacion"})
            t = _Tx(s)
            yield t
            t.flush()
            s.commit()
        except Exception:
            s.rollback()
            raise
        finally:
            s.close()
    else:
        with _lock:
            t = _Tx(None)
            yield t
            t.flush()


def peek_folio(prefix=FOLIO_PREFIX):
    """Siguiente folio SIN consumirlo (solo para mostrarlo en pantalla)."""
    if db_enabled():
        from sqlalchemy import text
        s = _orm.get_consig_session()
        try:
            n = s.execute(text("SELECT n FROM consignacion_folios WHERE prefix = :p"), {"p": prefix}).scalar()
        finally:
            s.close()
    else:
        n = _json_read(_data_dir / "folios.json", {}).get(prefix)
    return f"{prefix}-{str(int(n or 0) + 1).zfill(FOLIO_WIDTH)}"


# ══════════════════════════════════════════════════════════════════
#  Funciones usadas por el reporte de Job (app.py)
# ══════════════════════════════════════════════════════════════════
def _tag_items(order):
    return [{**i, "order_number": order.get("order_number", ""), "origen": ORIGEN}
            for i in (order.get("items") or [])]


def reassign_items_for_job(job_number, pool=None):
    """Items de reasignación de consignación para un Job, ya etiquetados con
    order_number y origen. Con pool (reporte multi-job) filtra en memoria; si no,
    filtra en SQL igual que reassign_items_for_job() de Stock."""
    job = (job_number or "").upper()
    if pool is None and db_enabled():
        from sqlalchemy import text
        s = _orm.get_consig_session()
        try:
            rows = (s.query(_orm.ConsigReasignacion)
                    .filter(text("EXISTS (SELECT 1 FROM jsonb_array_elements(consignacion_reasignaciones.data->'items') item "
                                 "WHERE UPPER(item->>'job') = :job)"))
                    .params(job=job).all())
            pool = [r.data for r in rows]
        finally:
            s.close()
    elif pool is None:
        pool = load("orders")
    return [i for o in pool for i in _tag_items(o) if (i.get("job") or "").upper() == job]


def recovery_for_job(job_number, pool=None):
    job = (job_number or "").upper()
    if pool is None and db_enabled():
        from sqlalchemy import func
        s = _orm.get_consig_session()
        try:
            pool = [r.data for r in s.query(_orm.ConsigRecuperacion)
                    .filter(func.upper(_orm.ConsigRecuperacion.job) == job).all()]
        finally:
            s.close()
    elif pool is None:
        pool = load("recovery")
    return [{**r, "origen": ORIGEN} for r in pool if (r.get("job") or "").upper() == job]


def export_rows(name):
    """(headers, rows) para /api/backup/<name>. Usa la unión de llaves de todos
    los registros (no solo las del primero) para no perder columnas."""
    if name == "consignacion":
        data = load("items")
    elif name == "consig-reassign":
        data = [{"order_number": o.get("order_number", ""), "created_at": o.get("created_at", ""), **i}
                for o in load("orders") for i in (o.get("items") or [])]
    elif name == "consig-recovery":
        data = load("recovery")
    else:
        return None
    if not data:
        return [], []
    headers = []
    for r in data:
        for k in r.keys():
            if k not in headers:
                headers.append(k)
    return headers, [[str(r.get(h, "")) for h in headers] for r in data]


# ══════════════════════════════════════════════════════════════════
#  Rutas — Inventario
# ══════════════════════════════════════════════════════════════════
def _filter_q(records, q, fields=("part_number", "manufacturer", "description", "label_code")):
    q = (q or "").lower()
    if not q:
        return records
    return [r for r in records if any(q in str(r.get(f, "") or "").lower() for f in fields)]


@bp.route("/api/consignacion", methods=["GET"])
def api_list():
    try:
        records = _filter_q(load("items"), request.args.get("q", ""))
        return jsonify({"records": records, "total": len(records), "db": db_enabled()})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def _recovery_record(item, qty, cost, job, stock_id):
    ts = _now()
    return {
        "id": f"CRCV-{ts.strftime('%Y%m%d%H%M%S%f')}",
        "manufacturer": item.get("manufacturer", ""), "part_number": item.get("part_number", ""),
        "description": item.get("description", ""),
        "last_cost": cost, "quantity": qty,
        "unit": item.get("unit", ""), "section": item.get("section", ""), "box": item.get("box", ""),
        "label_code": item.get("label_code", ""),
        "job": job, "total_value": round(qty * cost * -1, 2),
        "consig_id": stock_id, "created_at": ts.isoformat(),
    }


@bp.route("/api/consignacion", methods=["POST"])
def api_ingreso():
    """Ingreso de material (nuevo o existente). Mismo comportamiento que
    POST /api/stock: suma la cantidad nueva y, si trae Job de origen, registra
    una recuperación de costo para ese Job."""
    try:
        data = request.get_json() or {}
        mfr = str(data.get("manufacturer", "")).strip().upper()
        pnum = str(data.get("part_number", "")).strip().upper()
        if not mfr or not pnum:
            return jsonify({"error": "Fabricante y No. Parte son requeridos"}), 400
        new_qty = int(data.get("new_quantity", 0) or 0)
        if new_qty < 0:
            return jsonify({"error": "Los nuevos ingresos no pueden ser negativos"}), 400
        new_cost = float(data.get("last_cost", 0) or 0)
        recovery_job = str(data.get("recovery_job", "")).strip()
        unit = str(data.get("unit", "Pieza"))
        description = str(data.get("description", "")).strip()
        section = str(data.get("section", "")).strip()
        box = str(data.get("box", "")).strip()
        label_code = str(data.get("label_code", "")).strip().upper()
        ts = _now().isoformat()

        with tx() as t:
            records = t.load("items")
            existing = next((r for r in records
                             if r.get("manufacturer", "").upper() == mfr
                             and r.get("part_number", "").upper() == pnum), None)
            if existing:
                prev_qty = int(existing.get("quantity", 0) or 0)
                existing["quantity"] = prev_qty + new_qty
                if new_cost > 0:
                    existing["last_cost"] = new_cost
                if description and not existing.get("description"):
                    existing["description"] = description
                if section:      existing["section"] = section
                if box:          existing["box"] = box
                if recovery_job: existing["recovery_job"] = recovery_job
                if label_code:   existing["label_code"] = label_code
                existing["updated_at"] = ts
                item, action = existing, "updated"
            else:
                prev_qty = 0
                item = {
                    "id": f"CSG-{_now().strftime('%Y%m%d%H%M%S%f')}",
                    "manufacturer": mfr, "part_number": pnum, "description": description,
                    "last_cost": new_cost, "quantity": new_qty, "unit": unit,
                    "section": section, "box": box, "label_code": label_code,
                    "recovery_job": recovery_job, "created_at": ts,
                }
                records.append(item)
                action = "created"
            t.save("items", records)
            if new_qty > 0 and recovery_job:
                cost = new_cost if new_cost > 0 else float(item.get("last_cost", 0) or 0)
                recs = t.load("recovery")
                recs.append(_recovery_record(item, new_qty, cost, recovery_job, item["id"]))
                t.save("recovery", recs)
        return jsonify({"ok": True, "action": action, "record": item, "new_qty": new_qty, "prev_qty": prev_qty})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@bp.route("/api/consignacion/<item_id>", methods=["PUT"])
def api_update(item_id):
    try:
        data = request.get_json() or {}
        with tx() as t:
            records = t.load("items")
            rec = next((r for r in records if r.get("id") == item_id), None)
            if not rec:
                return jsonify({"error": "Item no encontrado"}), 404
            for k in ["manufacturer", "part_number", "description", "unit", "section", "box", "recovery_job", "label_code"]:
                if k in data:
                    rec[k] = str(data[k]).strip()
            for k in ("manufacturer", "part_number", "label_code"):
                rec[k] = (rec.get(k) or "").upper()
            if "last_cost" in data: rec["last_cost"] = float(data["last_cost"])
            if "quantity" in data:  rec["quantity"] = int(data["quantity"])
            rec["updated_at"] = _now().isoformat()
            t.save("items", records)
        return jsonify({"ok": True, "record": rec})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@bp.route("/api/consignacion/<item_id>", methods=["DELETE"])
def api_delete(item_id):
    if not _is_admin():
        return jsonify({"error": "Sin permiso"}), 403
    try:
        with tx() as t:
            records = t.load("items")
            new = [r for r in records if r.get("id") != item_id]
            if len(new) == len(records):
                return jsonify({"error": "No encontrado"}), 404
            t.save("items", new)
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def _norm_hdr(v):
    v = unicodedata.normalize("NFKD", str(v)).encode("ascii", "ignore").decode()
    return " ".join(v.strip().upper().split())


@bp.route("/api/consignacion/import", methods=["POST"])
def api_import():
    """Mismo formato y reglas que /api/stock/import (rev09), incluidos los
    encabezados en inglés del propio respaldo Excel, para poder re-importarlo."""
    try:
        f = request.files.get("file")
        if not f:
            return jsonify({"error": "No se recibio archivo"}), 400
        mode = request.form.get("mode", "append")
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
        ws = wb.active
        first = next(ws.iter_rows(min_row=1, max_row=1), None) or []
        headers = {_norm_hdr(c.value): c.column - 1 for c in first if c.value}

        def col(*aliases):
            for a in aliases:
                if a in headers:
                    return headers[a]
            return None
        ci_mfr = col("FABRICANTE", "MANUFACTURER", "MARCA")
        ci_pnum = col("NUMERO DE PARTE", "PART NUMBER", "PART_NUMBER", "NO. PARTE")
        ci_desc = col("DESCRIPCION", "DESCRIPTION", "DESC")
        ci_cost = col("ULTIMO COSTO", "LAST COST", "LAST_COST", "COSTO", "COST")
        ci_qty = col("EXISTENCIA", "QUANTITY", "CANTIDAD", "QTY")
        ci_unit = col("UNIDAD", "UNIT")
        ci_sec = col("SECCION", "SECTION")
        ci_box = col("CAJA", "BOX")
        ci_rec = col("RECUPERACION", "RECOVERY", "RECOVERY_JOB")
        ci_label = col("ETIQUETA", "LABEL", "LABEL_CODE", "QR", "CODIGO DE BARRAS", "BARCODE", "COD. ETIQUETA")
        missing = [n for n, c in (("FABRICANTE", ci_mfr), ("NUMERO DE PARTE", ci_pnum), ("EXISTENCIA", ci_qty)) if c is None]
        if missing:
            return jsonify({"error": "Faltan columnas requeridas en la primera fila de la hoja activa: " + ", ".join(missing)}), 400

        def cell(row, ci):
            if ci is None or ci >= len(row) or row[ci] is None:
                return None
            v = str(row[ci]).strip()
            return v or None

        now = _now()
        id_prefix = f"CSG-imp-{now.strftime('%Y%m%d%H%M%S')}"
        imported = created = updated = 0
        with tx() as t:
            records = t.load("items") if mode == "append" else []
            index = {(r.get("part_number", ""), r.get("manufacturer", "")): r for r in records}
            for row in ws.iter_rows(min_row=2, values_only=True):
                pnum = (cell(row, ci_pnum) or "").upper()
                if not pnum or pnum in ("NONE", "#N/A"):
                    continue
                mfr = (cell(row, ci_mfr) or "").upper()
                try: cost = float(cell(row, ci_cost)) if cell(row, ci_cost) else 0.0
                except Exception: cost = 0.0
                try: qty = int(float(cell(row, ci_qty))) if cell(row, ci_qty) else 0
                except Exception: qty = 0
                label = (cell(row, ci_label) or "").upper()
                desc, unit = cell(row, ci_desc) or "", cell(row, ci_unit) or ""
                section, box = cell(row, ci_sec) or "", cell(row, ci_box) or ""
                existing = index.get((pnum, mfr))
                if existing:
                    existing["quantity"] = qty
                    existing["last_cost"] = cost
                    if label:   existing["label_code"] = label
                    if unit:    existing["unit"] = unit
                    if section: existing["section"] = section
                    if box:     existing["box"] = box
                    if desc and not existing.get("description"): existing["description"] = desc
                    existing["updated_at"] = now.isoformat()
                    updated += 1
                else:
                    rec = {
                        "id": f"{id_prefix}-{created}", "manufacturer": mfr, "part_number": pnum,
                        "description": desc, "last_cost": cost, "quantity": qty,
                        "unit": unit or "Pieza", "section": section, "box": box,
                        "recovery_job": cell(row, ci_rec) or "", "label_code": label,
                        "created_at": now.isoformat(),
                    }
                    records.append(rec)
                    index[(pnum, mfr)] = rec
                    created += 1
                imported += 1
            t.save("items", records)
        return jsonify({"ok": True, "imported": imported, "created": created, "updated": updated, "total": len(records)})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ══════════════════════════════════════════════════════════════════
#  Rutas — Reasignaciones
# ══════════════════════════════════════════════════════════════════
@bp.route("/api/consignacion/reasignaciones", methods=["GET"])
def api_orders():
    try:
        job = request.args.get("job", "").strip().upper()
        orders = load("orders")
        if job:
            orders = [o for o in orders if any((i.get("job") or "").upper() == job for i in o.get("items", []))]
        return jsonify({"orders": orders, "total": len(orders), "next_number": peek_folio()})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@bp.route("/api/consignacion/reasignaciones", methods=["POST"])
def api_create_order():
    try:
        data = request.get_json() or {}
        is_new = data.get("is_new", True)
        items = data.get("items") or []
        if not items:
            return jsonify({"error": "Sin items"}), 400
        now = _now().isoformat()
        with tx() as t:
            orders = t.load("orders")
            records = t.load("items")
            # Validar TODO antes de modificar nada (cantidades acumuladas por material).
            pedido = {}
            for it in items:
                key = (str(it.get("part_number", "")).strip().upper(), str(it.get("manufacturer", "")).strip().upper())
                qty = int(it.get("quantity", 0) or 0)
                if qty < 1:
                    return jsonify({"error": f"Cantidad inválida para {key[0]}"}), 400
                if not str(it.get("job", "")).strip():
                    return jsonify({"error": f"Falta el Job destino para {key[0]}"}), 400
                pedido[key] = pedido.get(key, 0) + qty
            by_key = {(r.get("part_number", ""), r.get("manufacturer", "")): r for r in records}
            for key, qty in pedido.items():
                stk = by_key.get(key)
                if not stk:
                    return jsonify({"error": f"{key[0]} ({key[1]}) no existe en Consignación"}), 400
                if int(stk.get("quantity", 0) or 0) < qty:
                    return jsonify({"error": f"Existencia insuficiente de {key[0]}: hay {stk.get('quantity', 0)}, se piden {qty}"}), 400

            if is_new:
                order_number = t.next_folio()
                order = {"order_number": order_number, "created_at": now, "items": [],
                         "created_by": _current_user(), "origen": "Reasignación manual"}
                orders.append(order)
            else:
                order_number = str(data.get("order_number", "")).strip().upper()
                order = next((o for o in orders if o["order_number"] == order_number), None)
                if not order:
                    return jsonify({"error": "Orden no encontrada"}), 404

            for it in items:
                pnum = str(it.get("part_number", "")).strip().upper()
                mfr = str(it.get("manufacturer", "")).strip().upper()
                qty = int(it.get("quantity", 0))
                cost = float(it.get("unit_cost", 0) or 0)
                stk = by_key[(pnum, mfr)]
                stk["quantity"] = int(stk.get("quantity", 0) or 0) - qty
                stk["updated_at"] = now
                order["items"].append({
                    "part_number": pnum, "manufacturer": mfr,
                    "description": str(it.get("description", "")).strip(),
                    "label_code": str(it.get("label_code") or stk.get("label_code", "")).strip().upper(),
                    "job": str(it.get("job", "")).strip().upper(),
                    "unit_cost": cost, "quantity": qty, "total_cost": round(cost * qty, 2),
                    "added_at": now, "added_by": _current_user(),
                })
            order["updated_at"] = now
            t.save("orders", orders)
            t.save("items", records)
        return jsonify({"ok": True, "order_number": order_number, "order": order})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@bp.route("/api/consignacion/reasignaciones/<job_number>/total")
def api_order_total(job_number):
    try:
        items = reassign_items_for_job(job_number)
        return jsonify({"job": job_number, "total": round(sum(float(i.get("total_cost", 0)) for i in items), 2),
                        "items": items})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@bp.route("/api/consignacion/reasignaciones/orden/<order_number>", methods=["DELETE"])
def api_delete_order(order_number):
    if not _is_admin():
        return jsonify({"error": "Sin permiso"}), 403
    try:
        # Borrar la orden y regresar su material a Consignación en la MISMA transacción.
        with tx() as t:
            orders = t.load("orders")
            order = next((o for o in orders if o.get("order_number") == order_number.upper()), None)
            if not order:
                return jsonify({"error": "Orden no encontrada"}), 404
            records = t.load("items")
            devuelto = []
            now = _now()
            for n, it in enumerate(order.get("items") or []):
                pnum, mfr = str(it.get("part_number", "")).upper(), str(it.get("manufacturer", "")).upper()
                qty = int(float(it.get("quantity") or 0))
                if qty <= 0: continue
                rec = next((r for r in records if r.get("part_number", "") == pnum and r.get("manufacturer", "") == mfr), None)
                if rec:
                    rec["quantity"] = int(rec.get("quantity") or 0) + qty; rec["updated_at"] = now.isoformat(); accion = "sumado"
                else:
                    records.append({"id": f"CSG-dev-{now.strftime('%Y%m%d%H%M%S%f')}-{n}", "manufacturer": mfr, "part_number": pnum,
                                    "description": it.get("description", ""), "label_code": it.get("label_code", ""),
                                    "last_cost": float(it.get("unit_cost") or 0), "quantity": qty, "unit": "Pieza",
                                    "section": "", "box": "", "recovery_job": "", "created_at": now.isoformat()})
                    accion = "re-creado"
                devuelto.append({"part_number": pnum, "manufacturer": mfr, "quantity": qty, "accion": accion})
            t.save("items", records)
            t.save("orders", [o for o in orders if o is not order])
        return jsonify({"ok": True, "devuelto": devuelto})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def _h(v):
    import html
    return html.escape(str(v if v is not None else ""))


@bp.route("/api/consignacion/reasignaciones/orden/<order_number>/pdf")
def api_order_pdf(order_number):
    try:
        order = next((o for o in load("orders") if o["order_number"] == order_number.upper()), None)
        if not order:
            return jsonify({"error": "Orden no encontrada"}), 404
        rows, total = [], 0.0
        for it in order.get("items", []):
            t = float(it.get("total_cost", 0) or 0)
            total += t
            rows.append(f"""<tr><td>{_h(it.get('part_number'))}</td><td>{_h(it.get('manufacturer'))}</td>
              <td>{_h(it.get('description'))}</td><td>{_h(it.get('job'))}</td>
              <td style="text-align:right">{_h(it.get('quantity', 0))}</td>
              <td style="text-align:right">${float(it.get('unit_cost', 0) or 0):,.2f}</td>
              <td style="text-align:right">${t:,.2f}</td>
              <td>{_h(it.get('added_by') or order.get('created_by') or '—')}</td></tr>""")
        html = f"""<!DOCTYPE html><html><head><meta charset="UTF-8"><title>{_h(order['order_number'])}</title>
        <style>
          body{{font-family:Arial,sans-serif;font-size:11px;color:#222;margin:30px}}
          h1{{font-size:18px;color:#c8102e;margin-bottom:4px}}
          .tag{{display:inline-block;font-size:10px;font-weight:bold;letter-spacing:1px;color:#7a4b00;background:#fff3cd;border:1px solid #e0b252;border-radius:4px;padding:2px 8px;margin-bottom:8px}}
          .sub{{font-size:11px;color:#666;margin-bottom:20px}}
          table{{width:100%;border-collapse:collapse;margin-top:16px}}
          th{{background:#1f3864;color:#fff;padding:7px 10px;text-align:left;font-size:10px;text-transform:uppercase}}
          td{{padding:7px 10px;border-bottom:1px solid #e0e0e0;font-size:11px}}
          tr:nth-child(even){{background:#f5f5f5}}
          .total{{text-align:right;font-weight:bold;font-size:13px;color:#1a7a1a;margin-top:12px}}
          .footer{{margin-top:30px;font-size:10px;color:#999;border-top:1px solid #ddd;padding-top:8px}}
        </style></head><body>
        <div class="tag">MATERIAL EN CONSIGNACIÓN</div>
        <h1>Orden de Reasignación: {_h(order['order_number'])}</h1>
        <div class="sub">Fecha: {_h(order.get('created_at', '')[:10])} &nbsp;|&nbsp; Generada por: <b>{_h(order.get('created_by') or '— (orden anterior al registro de usuario)')}</b> &nbsp;|&nbsp; Persico México</div>
        <table><tr><th>No. Parte</th><th>Fabricante</th><th>Descripción</th><th>Job</th>
          <th style="text-align:right">Cant.</th><th style="text-align:right">Costo Unit.</th><th style="text-align:right">Total USD</th><th>Agregó</th></tr>
        {''.join(rows)}
        </table>
        <div class="total">Total: ${total:,.2f} USD</div>
        <div class="footer">Generado por Persico Suite · {_now().strftime('%Y-%m-%d %H:%M')}</div>
        </body></html>"""
        return Response(html, mimetype="text/html",
                        headers={"Content-Disposition": f"inline;filename={order['order_number']}.html"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ══════════════════════════════════════════════════════════════════
#  Rutas — Recuperaciones de consignación
# ══════════════════════════════════════════════════════════════════
@bp.route("/api/consignacion/recuperaciones", methods=["GET"])
def api_recovery():
    try:
        job = request.args.get("job", "").strip().upper()
        records = load("recovery")
        if job:
            records = [r for r in records if (r.get("job") or "").upper() == job]
        records = _filter_q(records, request.args.get("q", ""),
                            ("part_number", "manufacturer", "description", "label_code"))
        return jsonify({"records": records, "total": len(records)})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@bp.route("/api/consignacion/recuperaciones/<rec_id>", methods=["DELETE"])
def api_delete_recovery(rec_id):
    if not _is_admin():
        return jsonify({"error": "Sin permiso"}), 403
    try:
        with tx() as t:
            recs = t.load("recovery")
            new = [r for r in recs if r.get("id") != rec_id]
            if len(new) == len(recs):
                return jsonify({"error": "No encontrado"}), 404
            t.save("recovery", new)
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
